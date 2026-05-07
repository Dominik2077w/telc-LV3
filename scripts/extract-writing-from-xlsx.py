#!/usr/bin/env python3
"""Extract writing-topic flashcards from the German/Chinese topic workbook."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCE_SHEET = "德中题目对照"


def clean_text(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_topic_label(german: str) -> tuple[str, str]:
    match = re.match(r"^(\d+)([A-Z])\b", german)
    if not match:
        return german, ""
    return match.group(1), match.group(2)


def build_sections(input_path: Path) -> list[dict]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook[SOURCE_SHEET]
    rows = list(worksheet.iter_rows(values_only=True))
    header = [clean_text(cell) for cell in rows[0]]

    try:
        german_index = header.index("德语题目")
        chinese_index = header.index("中文对照")
    except ValueError as error:
        raise ValueError("Missing required columns: 德语题目, 中文对照") from error

    sections: dict[str, dict] = {}

    for source_row_number, row in enumerate(rows[1:], start=2):
        german = clean_text(row[german_index])
        chinese = clean_text(row[chinese_index])
        if not german:
            continue

        teil, letter = parse_topic_label(german)
        section = sections.setdefault(
            teil,
            {
                "teil": int(teil) if teil.isdigit() else teil,
                "title": f"Schreiben Thema {teil}",
                "sourceText": "德中写作题目对照.xlsx",
                "mode": "writing-card",
                "questions": [],
            },
        )

        section["questions"].append(
            {
                "number": letter or str(len(section["questions"]) + 1),
                "text": german,
                "german": german,
                "chinese": chinese or "暂无中文释义",
                "answer": "",
                "raw": german,
                "source": "writing-xlsx",
                "sourceRow": source_row_number,
                "paragraph": letter or str(len(section["questions"]) + 1),
                "firstSentence": german,
                "summary": chinese or "暂无中文释义",
            }
        )

    def sort_key(value: str) -> tuple[int, str]:
        return (int(value), "") if value.isdigit() else (10_000, value)

    return [sections[key] for key in sorted(sections, key=sort_key)]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path, help="Path to writing topic workbook")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/writing_questions.js"),
        help="Output JavaScript data file",
    )
    args = parser.parse_args()

    sections = build_sections(args.input)
    payload = json.dumps(sections, ensure_ascii=False, indent=2)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(f"window.WRITING_QUESTIONS = {payload};\n", encoding="utf-8")

    question_count = sum(len(section["questions"]) for section in sections)
    print(f"Wrote {len(sections)} writing sections and {question_count} questions to {args.output}")


if __name__ == "__main__":
    main()
